#!/usr/bin/env python3
"""Assemble rendered slide PNGs into a PPTX (with speaker notes) and a PDF.

    python3 tools/make_deck.py course/video/lessons/W01-L02.json slides/dir out/dir
"""
import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Emu

spec_path, slide_dir, out_dir = map(Path, sys.argv[1:4])
spec = json.loads(spec_path.read_text())
pngs = sorted(slide_dir.glob("slide-*.png"))
assert len(pngs) == len(spec["scenes"]), (len(pngs), len(spec["scenes"]))
out_dir.mkdir(parents=True, exist_ok=True)
name = f"Becoming-Doom-{spec['id']}-slides"

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
for png, sc in zip(pngs, spec["scenes"]):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
    s.notes_slide.notes_text_frame.text = "\n\n".join(sc.get("say", []))
prs.save(out_dir / f"{name}.pptx")

ims = [Image.open(p).convert("RGB") for p in pngs]
ims[0].save(out_dir / f"{name}.pdf", save_all=True, append_images=ims[1:], resolution=144)

# Contact sheet for quick review: three columns of thumbnails.
w, h = 640, 360
rows = (len(ims) + 2) // 3
sheet = Image.new("RGB", (w * 3, h * rows), (6, 6, 6))
for i, im in enumerate(ims):
    sheet.paste(im.resize((w, h)), ((i % 3) * w, (i // 3) * h))
sheet.save(out_dir / f"{name}-overview.jpg", quality=85)
print(out_dir / name)
