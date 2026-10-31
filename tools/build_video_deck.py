#!/usr/bin/env python3
"""Build a Becoming Doom lesson video deck from a lesson script.

Usage: python3 tools/build_video_deck.py course/video/W01-L01.md TEMPLATE.pptx OUT.pptx

The template supplies the theme, masters and fonts (DM Sans, Bebas Neue).
Every slide is drawn on the template's "Black (Blank)" layout so the
Becoming Doom identity (basalt black, ivory type, restrained gold) stays
consistent. The script's "Say" text becomes the speaker notes.
"""
import copy
import re
import sys
from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

INK = RGBColor(0x0B, 0x0B, 0x0B)
IVORY = RGBColor(0xEE, 0xEA, 0xE0)
STONE = RGBColor(0x8A, 0x86, 0x7E)
GRAPHITE = RGBColor(0x2A, 0x2A, 0x2A)
GOLD = RGBColor(0xC9, 0xA2, 0x4A)

W, H = 12192000, 6858000
M = 914400  # outer margin, matching the template grid
SANS = "DM Sans"
DISPLAY = "Bebas Neue"


def parse(path):
    text = Path(path).read_text()
    head, _, rest = text.partition("\n---\n")
    lesson_id = re.search(r"^# (W\d\d-L\d\d)", head, re.M).group(1)
    title = re.search(r"^# W\d\d-L\d\d: (.+)$", head, re.M).group(1)
    body = rest.split("\n---\n")[0]
    slides = []
    for block in re.split(r"^## Slide ", body, flags=re.M)[1:]:
        header, _, content = block.partition("\n")
        num, kind, timing = [p.strip() for p in header.split("|")]
        fields = dict(re.findall(r"^\*\*(On screen|Kicker|Visual|Items):\*\* (.+)$", content, re.M))
        say = content.split("**Say:**", 1)[1].strip() if "**Say:**" in content else ""
        items = []
        if "Items" in fields:
            for chunk in fields["Items"].split(" | "):
                label, _, detail = chunk.partition(" :: ")
                items.append((label.strip(), detail.strip()))
        slides.append(dict(num=int(num), kind=kind, timing=timing, screen=fields.get("On screen", ""),
                           kicker=fields.get("Kicker", ""), visual=fields.get("Visual"), items=items, say=say))
    return lesson_id, title, slides


def clear_slides(prs):
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.get(qn("r:id")))
        ids.remove(sld)


def layout(prs, name):
    return next(l for l in prs.slide_layouts if l.name == name)


def add_text(slide, x, y, w, h, runs, size, font=SANS, color=IVORY, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spacing=None, line=None):
    """runs: plain string, with **word** marking gold emphasis."""
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line:
            p.line_spacing = line
        for j, part in enumerate(re.split(r"\*\*(.+?)\*\*", para)):
            if not part:
                continue
            r = p.add_run()
            r.text = part
            f = r.font
            f.name = font
            f.size = Pt(size)
            f.bold = bold
            f.color.rgb = GOLD if j % 2 else color
            if spacing is not None:
                r._r.get_or_add_rPr().set("spc", str(spacing))
    return box


def rect(slide, x, y, w, h, color, alpha=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    s.line.fill.background()
    s.fill.solid()
    s.fill.fore_color.rgb = color
    if alpha is not None:
        srgb = s.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        etree.SubElement(srgb, qn("a:alpha")).set("val", str(int(alpha * 1000)))
    return s


def gradient(slide, x, y, w, h, stops, angle):
    """stops: [(position 0-100, alpha 0-100)] of black. angle in degrees."""
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    s.line.fill.background()
    spPr = s.fill._xPr
    for tag in ("a:solidFill", "a:noFill", "a:gradFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    grad = etree.Element(qn("a:gradFill"), rotWithShape="1")
    lst = etree.SubElement(grad, qn("a:gsLst"))
    for pos, a in stops:
        gs = etree.SubElement(lst, qn("a:gs"), pos=str(pos * 1000))
        c = etree.SubElement(gs, qn("a:srgbClr"), val="0B0B0B")
        etree.SubElement(c, qn("a:alpha"), val=str(a * 1000))
    etree.SubElement(grad, qn("a:lin"), ang=str(angle * 60000), scaled="0")
    spPr.insert(1, grad)
    return s


def placeholder_image(path, label, size):
    img = Image.new("RGB", size, (20, 20, 20))
    d = ImageDraw.Draw(img)
    for i in range(0, size[0] + size[1], 40):
        d.line([(i, 0), (i - size[1], size[1])], fill=(28, 28, 28), width=2)
    d.text((40, size[1] - 60), f"IMAGE: {label}", fill=(201, 162, 74))
    img.save(path)


def picture(slide, assets, name, x, y, w, h, scratch):
    """Place an image cropped to cover the box."""
    src = assets / f"{name}.png"
    if not src.exists():
        src = scratch / f"placeholder-{name}-{w}x{h}.png"
        placeholder_image(src, name, (1600, int(1600 * h / w)))
    with Image.open(src) as im:
        iw, ih = im.size
    pic = slide.shapes.add_picture(str(src), Emu(x), Emu(y), Emu(w), Emu(h))
    box_r, img_r = w / h, iw / ih
    if img_r > box_r:
        cut = (1 - box_r / img_r) / 2
        pic.crop_left = pic.crop_right = cut
    else:
        cut = (1 - img_r / box_r) / 2
        pic.crop_top = pic.crop_bottom = cut
    return pic


def chrome(slide, lesson_id, num, total, light_rule=True):
    add_text(slide, M, 420000, 4000000, 300000, lesson_id, 10, color=STONE, spacing=300)
    add_text(slide, W - M - 1500000, 420000, 1500000, 300000, f"{num:02d} / {total:02d}", 10,
             color=STONE, align=PP_ALIGN.RIGHT, spacing=300)
    add_text(slide, W - M - 3000000, H - 620000, 3000000, 300000, "BECOMING DOOM", 12, font=DISPLAY,
             color=STONE, align=PP_ALIGN.RIGHT, spacing=600)
    if light_rule:
        rect(slide, M, H - 520000, 420000, 18000, GOLD)


def ghost_number(slide, num):
    add_text(slide, W - M - 4600000, 900000, 4600000, 5200000, f"{num:02d}", 300, font=DISPLAY,
             color=RGBColor(0x17, 0x17, 0x17), align=PP_ALIGN.RIGHT)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = f"[{s['timing']}]\n\n" + s["say"].replace("**", "")


def build(script, template, out):
    lesson_id, title, slides = parse(script)
    assets = Path(script).parent / "assets" / lesson_id
    scratch = Path(out).parent
    scratch.mkdir(parents=True, exist_ok=True)
    prs = Presentation(template)
    clear_slides(prs)
    blank = layout(prs, "Black (Blank)")
    total = len(slides)

    for s in slides:
        sl = prs.slides.add_slide(blank)
        k = s["kind"]

        if k in ("title", "close", "image-full"):
            picture(sl, assets, s["visual"], 0, 0, W, H, scratch)
            gradient(sl, 0, 0, W, H, [(0, 92), (45, 70), (100, 0)], 0)
            gradient(sl, 0, H // 2, W, H // 2, [(0, 0), (100, 80)], 90)
            if k == "title":
                add_text(sl, M, 2050000, 6200000, 400000, s["kicker"].upper(), 12, color=GOLD, spacing=400)
                add_text(sl, M, 2500000, 6400000, 2600000, s["screen"].upper(), 96, font=DISPLAY,
                         line=0.85)
            elif k == "close":
                add_text(sl, M, 1900000, 5800000, 2800000, s["screen"], 36, line=1.1)
                rect(sl, M, 4950000, 420000, 18000, GOLD)
                add_text(sl, M, 5100000, 5800000, 400000, s["kicker"].upper(), 12, color=GOLD, spacing=400)
            else:
                add_text(sl, M, 4300000, 6000000, 1600000, s["screen"], 40, line=1.0,
                         anchor=MSO_ANCHOR.BOTTOM)
            chrome(sl, lesson_id, s["num"], total, light_rule=False)

        elif k == "hook":
            if s["visual"]:
                picture(sl, assets, s["visual"], 0, 0, W, H, scratch)
                rect(sl, 0, 0, W, H, INK, alpha=78)
            add_text(sl, M + 600000, 1500000, W - 2 * M - 1200000, 3800000, s["screen"], 48,
                     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line=1.05)
            chrome(sl, lesson_id, s["num"], total)

        elif k == "statement":
            ghost_number(sl, s["num"])
            rect(sl, M, 1900000, 420000, 18000, GOLD)
            add_text(sl, M, 2150000, 6300000, 3600000, s["screen"], 54, line=1.0)
            chrome(sl, lesson_id, s["num"], total, light_rule=False)

        elif k == "image-right":
            pw = 5410200
            picture(sl, assets, s["visual"], W - pw, 0, pw, H, scratch)
            gradient(sl, W - pw, 0, 1400000, H, [(0, 100), (100, 0)], 0)
            rect(sl, M, 2300000, 420000, 18000, GOLD)
            add_text(sl, M, 2550000, 5600000, 2600000, s["screen"], 54, line=1.0)
            chrome(sl, lesson_id, s["num"], total, light_rule=False)

        elif k in ("three", "mission"):
            add_text(sl, M, 1500000, 3600000, 2400000, s["screen"], 34, line=1.05)
            rect(sl, M, 1250000, 420000, 18000, GOLD)
            x0, top = 5200000, 1500000
            row = 1450000
            for i, (label, detail) in enumerate(s["items"]):
                y = top + i * row
                rect(sl, x0, y - 120000, W - M - x0, 9000, GRAPHITE)
                if k == "three":
                    add_text(sl, x0, y, 900000, 900000, f"{i + 1:02d}", 40, font=DISPLAY, color=GOLD)
                    add_text(sl, x0 + 1000000, y + 30000, W - M - x0 - 1000000, 400000, label, 22, bold=True)
                    add_text(sl, x0 + 1000000, y + 480000, W - M - x0 - 1000000, 700000, detail, 18,
                             color=STONE)
                else:
                    add_text(sl, x0, y + 30000, 1600000, 400000, label.upper(), 12, color=GOLD, spacing=300)
                    add_text(sl, x0 + 1800000, y, W - M - x0 - 1800000, 1100000, detail, 20)
            chrome(sl, lesson_id, s["num"], total, light_rule=False)

        elif k == "steps":
            add_text(sl, M, 1150000, 9000000, 700000, s["screen"], 34)
            cw, ch = (W - 2 * M - 600000) // 2, 1900000
            for i, (label, detail) in enumerate(s["items"]):
                cx = M + (i % 2) * (cw + 600000)
                cy = 2250000 + (i // 2) * ch
                rect(sl, cx, cy, cw, 9000, GRAPHITE)
                add_text(sl, cx, cy + 200000, 1300000, 1000000, f"{label}.", 60, font=DISPLAY, color=GOLD)
                add_text(sl, cx + 1350000, cy + 300000, cw - 1350000, 1300000, detail, 20, line=1.1)
            chrome(sl, lesson_id, s["num"], total)
        else:
            raise ValueError(f"unknown slide kind {k}")
        notes(sl, s)

    prs.core_properties.title = f"{lesson_id}: {title}"
    prs.core_properties.author = "Becoming Doom"
    prs.save(out)
    print(f"{out}: {total} slides")


if __name__ == "__main__":
    build(*sys.argv[1:4])
